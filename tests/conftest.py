"""
Shared fixtures for all bank migration tests.
"""
import os
import sys
import json
import csv
import tempfile
import shutil
from pathlib import Path
from unittest.mock import MagicMock

# Ensure src is importable
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest


# ---------------------------------------------------------------------------
# Temporary directory fixture
# ---------------------------------------------------------------------------

@pytest.fixture
def tmp_dir():
    """Provide a temporary directory that is cleaned up after the test."""
    d = tempfile.mkdtemp()
    yield d
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def sample_dir(tmp_dir):
    """Create a subdirectory for sample files."""
    d = os.path.join(tmp_dir, "samples")
    os.makedirs(d, exist_ok=True)
    return d


# ---------------------------------------------------------------------------
# Sample data fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def sample_record():
    """A single valid customer record."""
    return {
        "full_name": "Muhammad Tayyab",
        "dob": "1995-03-15",
        "account_number": "1234567890123456",
        "email": "tayyab@example.com",
        "phone": "03001234567",
        "address": "123 Main St, Lahore, Punjab, Pakistan",
        "balance": "50000.00",
    }


@pytest.fixture
def sample_records():
    """A list of valid customer records."""
    return [
        {
            "full_name": "Muhammad Tayyab",
            "dob": "1995-03-15",
            "account_number": "1234567890123456",
            "email": "tayyab@example.com",
            "phone": "03001234567",
            "address": "123 Main St, Lahore, Punjab, Pakistan",
            "balance": "50000.00",
        },
        {
            "full_name": "Ali Ahmed Khan",
            "dob": "1988-07-22",
            "account_number": "9876543210987654",
            "email": "ali.ahmed@test.org",
            "phone": "03219876543",
            "address": "456 Oak Ave, Karachi, Sindh 75500, Pakistan",
            "balance": "$125000.50",
        },
        {
            "full_name": "Sara",
            "dob": "2000-01-01",
            "account_number": "1111222233334444",
            "email": "sara@mail.com",
            "phone": "03331112222",
            "address": "789 Pine Rd, Islamabad, ICT, Pakistan",
            "balance": "₨75000",
        },
    ]


# ---------------------------------------------------------------------------
# File fixtures — create real files on disk for FormatDetector tests
# ---------------------------------------------------------------------------

@pytest.fixture
def csv_file(sample_dir, sample_records):
    """Create a real CSV file with sample records."""
    path = os.path.join(sample_dir, "test_data.csv")
    fieldnames = sample_records[0].keys()
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(sample_records)
    return path


@pytest.fixture
def json_file(sample_dir, sample_records):
    """Create a real JSON file with sample records."""
    path = os.path.join(sample_dir, "test_data.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(sample_records, f)
    return path


@pytest.fixture
def json_single_record(sample_dir):
    """Create a JSON file with a single record (dict, not list)."""
    path = os.path.join(sample_dir, "single_record.json")
    record = {"full_name": "Test User", "dob": "1990-01-01"}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(record, f)
    return path


@pytest.fixture
def xml_file(sample_dir):
    """Create a real XML file with sample records."""
    path = os.path.join(sample_dir, "test_data.xml")
    content = """<?xml version="1.0" encoding="UTF-8"?>
<records>
    <record>
        <full_name>Muhammad Tayyab</full_name>
        <dob>1995-03-15</dob>
        <account_number>1234567890123456</account_number>
        <email>tayyab@example.com</email>
        <phone>03001234567</phone>
    </record>
    <record>
        <full_name>Ali Ahmed</full_name>
        <dob>1988-07-22</dob>
        <account_number>9876543210987654</account_number>
        <email>ali@test.org</email>
        <phone>03219876543</phone>
    </record>
</records>"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


@pytest.fixture
def txt_file(sample_dir):
    """Create a real TXT file."""
    path = os.path.join(sample_dir, "test_data.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write("This is a plain text migration file.\nLine 2 of data.")
    return path


@pytest.fixture
def xlsx_file(sample_dir, sample_records):
    """Create a real XLSX file with sample records."""
    from openpyxl import Workbook
    path = os.path.join(sample_dir, "test_data.xlsx")
    wb = Workbook()
    ws = wb.active
    headers = list(sample_records[0].keys())
    ws.append(headers)
    for rec in sample_records:
        ws.append([rec.get(h, "") for h in headers])
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def docx_file(sample_dir):
    """Create a real DOCX file with a table of records."""
    from docx import Document
    path = os.path.join(sample_dir, "test_data.docx")
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    # Header row
    headers = ["full_name", "dob", "email"]
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    # Data rows
    table.rows[1].cells[0].text = "Muhammad Tayyab"
    table.rows[1].cells[1].text = "1995-03-15"
    table.rows[1].cells[2].text = "tayyab@example.com"
    table.rows[2].cells[0].text = "Ali Ahmed"
    table.rows[2].cells[1].text = "1988-07-22"
    table.rows[2].cells[2].text = "ali@test.org"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Component fixtures — pre-built instances for unit tests
# ---------------------------------------------------------------------------

@pytest.fixture
def parser():
    from src.parser import Parser
    return Parser()


@pytest.fixture
def validator():
    from src.validator import Validator
    return Validator()


@pytest.fixture
def validator_with_rules():
    from src.validator import Validator
    return Validator(rules={
        "dob": {"type": "date"},
        "email": {"type": "email"},
        "phone": {"type": "phone"},
        "account_number": {"min_length": 5},
        "full_name": {"required": True},
    })


@pytest.fixture
def masker():
    from src.security import SecurityMasker
    return SecurityMasker()


@pytest.fixture
def masker_with_audit():
    from src.security import SecurityMasker
    from src.audit_logger import AuditLogger
    audit = AuditLogger(migration_id="test_migration")
    return SecurityMasker(audit_logger=audit), audit


@pytest.fixture
def audit_logger(tmp_dir):
    from src.audit_logger import AuditLogger
    logger = AuditLogger(migration_id="test_audit")
    # Override log path to use tmp_dir
    logger._log_path = Path(tmp_dir) / "audit_test.jsonl"
    return logger


@pytest.fixture
def txn_manager():
    from src.transaction_rollback import TransactionManager
    return TransactionManager()


@pytest.fixture
def rules_engine():
    from src.rules_engine import RulesEngine, build_standard_rules
    return RulesEngine(build_standard_rules())


@pytest.fixture
def bank_registry():
    from src.registry import BankRegistry
    return BankRegistry()


@pytest.fixture
def schema_mapper(bank_registry):
    from src.schema_mapper import SchemaMapper
    return SchemaMapper(registry=bank_registry)


@pytest.fixture
def canonical_store(tmp_dir):
    from src.canonical_store import CanonicalStore
    store = CanonicalStore(
        encryption_key="test-key-123",
        db_manager=MagicMock(),
    )
    return store


@pytest.fixture
def mock_components():
    """Provide all mocked components for isolated Transformer tests."""
    from src.validator import Validator
    from src.parser import Parser
    from src.schema_mapper import SchemaMapper
    from src.rules_engine import RulesEngine, build_standard_rules
    from src.security import SecurityMasker
    from src.audit_logger import AuditLogger
    from src.canonical_store import CanonicalStore
    from src.transaction_rollback import TransactionManager
    from src.registry import BankRegistry

    audit = AuditLogger(migration_id="test_pipeline")
    canonical = CanonicalStore(db_manager=MagicMock(), encryption_key="test-key")
    txn = TransactionManager()
    validator = Validator()
    parser = Parser()
    registry = BankRegistry()
    mapper = SchemaMapper(registry=registry)
    rules = RulesEngine(build_standard_rules())
    masker = SecurityMasker(audit_logger=audit)

    return {
        "audit": audit,
        "canonical": canonical,
        "txn": txn,
        "validator": validator,
        "parser": parser,
        "mapper": mapper,
        "rules": rules,
        "masker": masker,
    }
